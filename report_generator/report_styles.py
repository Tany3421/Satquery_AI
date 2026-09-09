import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    """Set background fill color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    """Set subtle horizontal borders on table with no vertical borders."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="none"/>
            <w:left w:val="none"/>
            <w:right w:val="none"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

class ReportStyleManager:
    def __init__(self, doc):
        self.doc = doc
        self.primary_hex = "1A365D"      # Deep Academic Navy
        self.secondary_hex = "2B6CB0"    # Rich Slate Blue
        self.accent_hex = "0D9488"       # Geospatial Teal
        self.dark_neutral_hex = "1E293B" # Slate 800
        self.light_bg_hex = "F8FAFC"     # Slate 50
        self.border_hex = "CBD5E1"       # Slate 300
        
        self.setup_page_layout()
        self.setup_base_styles()

    def setup_page_layout(self):
        for section in self.doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)
            section.page_width = Inches(8.5)
            section.page_height = Inches(11.0)
            
            # Header
            header = section.header
            hp = header.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            hrun = hp.add_run("SatQuery AI: Academic Project Report | Multimodal Geospatial Intelligence")
            hrun.font.name = "Calibri"
            hrun.font.size = Pt(8.5)
            hrun.font.color.rgb = RGBColor(100, 116, 139)
            
            # Footer
            footer = section.footer
            fp = footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            frun = fp.add_run("Department of Computer Science & Engineering | Academic Year 2025–2026")
            frun.font.name = "Calibri"
            frun.font.size = Pt(8.5)
            frun.font.color.rgb = RGBColor(100, 116, 139)

    def setup_base_styles(self):
        normal = self.doc.styles['Normal']
        normal.font.name = 'Calibri'
        normal.font.size = Pt(11)
        normal.font.color.rgb = RGBColor(30, 41, 59)
        normal.paragraph_format.line_spacing = 1.15
        normal.paragraph_format.space_after = Pt(6)

    def heading_1(self, text, space_before=20, space_after=8):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.bold = True
        run.font.size = Pt(17)
        run.font.color.rgb = RGBColor(26, 54, 93) # Navy
        
        # Add stylish bottom accent border
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(f'''
            <w:pBdr {nsdecls("w")}>
                <w:bottom w:val="single" w:sz="12" w:space="4" w:color="{self.secondary_hex}"/>
            </w:pBdr>
        ''')
        pPr.append(pBdr)
        return p

    def heading_2(self, text, space_before=14, space_after=6):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.bold = True
        run.font.size = Pt(13.5)
        run.font.color.rgb = RGBColor(43, 108, 176) # Slate Blue
        return p

    def heading_3(self, text, space_before=10, space_after=4):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.bold = True
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(13, 148, 136) # Teal
        return p

    def para(self, text, bold_prefix="", space_after=6, italic=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        
        if bold_prefix:
            rp = p.add_run(bold_prefix)
            rp.font.name = "Calibri"
            rp.font.bold = True
            rp.font.color.rgb = RGBColor(26, 54, 93)
            
        rt = p.add_run(text)
        rt.font.name = "Calibri"
        rt.font.size = Pt(10.5)
        rt.font.italic = italic
        rt.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def bullet(self, title, desc):
        p = self.doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        rt = p.add_run(title)
        rt.font.name = "Calibri"
        rt.font.bold = True
        rt.font.size = Pt(10.5)
        rt.font.color.rgb = RGBColor(26, 54, 93)
        
        rd = p.add_run(f": {desc}")
        rd.font.name = "Calibri"
        rd.font.size = Pt(10.5)
        rd.font.color.rgb = RGBColor(51, 65, 85)
        return p

    def callout(self, text, title="NOTE"):
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.5)
        
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F0FDF4" if "TAKEAWAY" in title or "KEY" in title else "F8FAFC")
        set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="{self.accent_hex}"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        
        rt = p.add_run(f"[{title}] ")
        rt.font.name = "Calibri"
        rt.font.bold = True
        rt.font.size = Pt(10)
        rt.font.color.rgb = RGBColor(13, 148, 136)
        
        rc = p.add_run(text)
        rc.font.name = "Calibri"
        rc.font.size = Pt(10)
        rc.font.color.rgb = RGBColor(30, 41, 59)
        
        ps = self.doc.add_paragraph()
        ps.paragraph_format.space_after = Pt(4)
        return tbl

    def diagram_box(self, diagram_text, caption=""):
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.5)
        
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F1F5F9")
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
                <w:left w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
                <w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
            </w:tcBorders>
        ''')
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        
        r = p.add_run(diagram_text.strip())
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(15, 23, 42)
        
        if caption:
            p_cap = self.doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(4)
            p_cap.paragraph_format.space_after = Pt(8)
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rcap = p_cap.add_run(caption)
            rcap.font.name = "Calibri"
            rcap.font.size = Pt(9.5)
            rcap.font.italic = True
            rcap.font.bold = True
            rcap.font.color.rgb = RGBColor(71, 85, 105)
        else:
            ps = self.doc.add_paragraph()
            ps.paragraph_format.space_after = Pt(4)

    def table(self, headers, rows, col_widths=None, caption=""):
        tbl = self.doc.add_table(rows=len(rows) + 1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        set_table_borders(tbl, color="CBD5E1")
        
        if col_widths:
            for i, w in enumerate(col_widths):
                tbl.columns[i].width = Inches(w)
                
        # Header Row
        for j, h in enumerate(headers):
            cell = tbl.cell(0, j)
            if col_widths:
                cell.width = Inches(col_widths[j])
            set_cell_background(cell, self.primary_hex)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h)
            r.font.name = "Calibri"
            r.font.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)
            
        # Data Rows
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = tbl.cell(i + 1, j)
                if col_widths:
                    cell.width = Inches(col_widths[j])
                if (i % 2) == 1:
                    set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=70, bottom=70, left=120, right=120)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                r = p.add_run(str(val))
                r.font.name = "Calibri"
                r.font.size = Pt(9.0)
                r.font.color.rgb = RGBColor(30, 41, 59)
                
        if caption:
            p_cap = self.doc.add_paragraph()
            p_cap.paragraph_format.space_before = Pt(4)
            p_cap.paragraph_format.space_after = Pt(8)
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rcap = p_cap.add_run(caption)
            rcap.font.name = "Calibri"
            rcap.font.size = Pt(9.5)
            rcap.font.italic = True
            rcap.font.bold = True
            rcap.font.color.rgb = RGBColor(71, 85, 105)
        else:
            ps = self.doc.add_paragraph()
            ps.paragraph_format.space_after = Pt(4)
            
        return tbl

    def page_break(self):
        self.doc.add_page_break()
